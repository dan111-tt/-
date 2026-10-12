from __future__ import annotations
import json, os, re, socket, ipaddress, uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import FastAPI, Depends, HTTPException, Security
from fastapi.security import APIKeyHeader
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field, HttpUrl

from app.research_rules import build_general_queries, country_name, draft_signal, extract_published_date, source_for_url

ROOT = Path(__file__).resolve().parent.parent
SOURCES = json.loads((ROOT / "config/sources.json").read_text(encoding="utf-8"))
SEARCH_TERMS = json.loads((ROOT / "config/search_terms.json").read_text(encoding="utf-8"))
OUT = Path(os.getenv("REPORT_DIR", "/tmp/information-collector-reports")); OUT.mkdir(parents=True, exist_ok=True)
API_KEY = os.getenv("API_KEY", "")
KEY_HEADER = APIKeyHeader(name="X-API-Key", auto_error=False)

app = FastAPI(title="Configurable Information Collector", version="0.3.0", description="Collect and verify public information using user-provided task parameters.")

def auth(key: str = Security(KEY_HEADER)):
    if not API_KEY or key != API_KEY:
        raise HTTPException(401, "Invalid API key")

class ResearchRequest(BaseModel):
    task_name: str = "资讯收集任务"
    start_date: date
    end_date: date
    regions: list[str] = Field(default_factory=list, description="User-provided countries, regions, markets, or geographic scope")
    industry_fields: list[str] = Field(default_factory=list, description="User-defined industries or fields")
    collection_scope: list[str] = Field(default_factory=list, description="User-defined subjects, questions, or information scope")
    source_domains: list[str] = Field(default_factory=list, description="Optional user-provided allowed domains")
    exclude_terms: list[str] = Field(default_factory=list)
    # Legacy fields remain accepted for existing regulatory-monitor actions.
    countries: list[str] = Field(default_factory=list)
    topics: list[str] = Field(default_factory=list)
    keywords: list[str] = Field(default_factory=list)
    exclude_drafts: bool = False
    urls: list[HttpUrl] = Field(default_factory=list, description="Optional known official/article URLs")
    source_urls: list[HttpUrl] = Field(default_factory=list, description="Optional user-uploaded source links")
    limit: int = Field(default=10, ge=1, le=30, description="Maximum candidates returned per region after discovery")

class ReportItem(BaseModel):
    region: str = ""
    country: str = ""
    industry: str = ""
    topic: str = "其他"
    title: str
    summary: str
    published_date: str = ""
    status: str = "待核验"
    approved: bool = False
    sources: list[HttpUrl] = Field(default_factory=list)

class ReportRequest(BaseModel):
    title: str = "资讯收集报告"
    period: str
    updated_at: str = Field(default_factory=lambda: date.today().isoformat())
    items: list[ReportItem] = Field(min_length=1)

@app.get("/health")
def health():
    return {"ok": True, "search_enabled": bool(os.getenv("TAVILY_API_KEY"))}

@app.get("/sources", dependencies=[Depends(auth)])
def sources(region: str | None = None, country: str | None = None):
    target = region or country  # country is a legacy alias
    return [s for s in SOURCES if not target or country_name(s["country"]) == country_name(target)]

def _host_matches(host: str, allowed_hosts: list[str]) -> bool:
    return any(host == h or host.endswith("."+h) for h in allowed_hosts)

def allowed(url: str, additional_hosts: list[str] | None = None, allow_any_public: bool = False) -> bool:
    host = (urlparse(url).hostname or "").lower().rstrip(".")
    if not host or urlparse(url).scheme != "https": return False
    if additional_hosts is not None:
        return _host_matches(host, additional_hosts) or allow_any_public
    return _host_matches(host,[s["host"] for s in SOURCES])

def _public_host(host: str):
    try:
        ips = [ipaddress.ip_address(x[4][0]) for x in socket.getaddrinfo(host, None)]
        return bool(ips) and all(ip.is_global for ip in ips)
    except Exception: return False

def fetch(url: str, additional_hosts: list[str] | None = None, allow_any_public: bool = False) -> dict | None:
    if not allowed(url, additional_hosts, allow_any_public): return None
    host = urlparse(url).hostname or ""
    if not _public_host(host): return None
    try:
        with httpx.Client(timeout=12, follow_redirects=True, headers={"User-Agent":"InformationCollector/0.3 (public information research)"}) as c:
            r = c.get(url)
            final_host=(urlparse(str(r.url)).hostname or "").lower()
            if not allowed(str(r.url), additional_hosts, allow_any_public) or not _public_host(final_host) or r.status_code != 200: return None
            soup = BeautifulSoup(r.text[:2_000_000], "html.parser")
            for x in soup(["style", "nav", "footer", "header", "noscript"]): x.decompose()
            for x in soup.find_all("script"):
                if x.get("type") != "application/ld+json": x.decompose()
            title = (soup.title.get_text(" ", strip=True) if soup.title else "").strip()
            published = ""
            for key, value in (("property", "article:published_time"), ("property", "og:published_time"), ("name", "pubdate"), ("name", "publishdate"), ("name", "date"), ("name", "dc.date"), ("itemprop", "datePublished")):
                meta = soup.find("meta", attrs={key:value})
                if meta and meta.get("content"):
                    published = meta["content"]; break
            structured = ""
            for tag in soup.find_all("script", attrs={"type":"application/ld+json"}):
                try:
                    data = json.loads(tag.string or tag.get_text())
                    stack = data if isinstance(data, list) else [data]
                    while stack:
                        obj = stack.pop()
                        if isinstance(obj, dict):
                            if obj.get("datePublished"): structured = str(obj["datePublished"]); break
                            stack.extend(v for v in obj.values() if isinstance(v, (dict,list)))
                        elif isinstance(obj, list): stack.extend(obj)
                    if structured: break
                except Exception: pass
            text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))[:18000]
            return {"url":str(r.url), "title":title, "published_date":published, "structured_date":structured, "text":text}
    except Exception: return None

def search_tavily(query: str, limit: int, domains: list[str] | None = None, allow_any_public: bool = False):
    key = os.getenv("TAVILY_API_KEY")
    if not key: return {"results":[],"error":"API key unavailable"}
    try:
        payload={"api_key":key,"query":query,"max_results":limit,"search_depth":"advanced","include_raw_content":False}
        if domains: payload["include_domains"] = domains[:20]
        r = httpx.post("https://api.tavily.com/search", json=payload, timeout=25)
        r.raise_for_status()
        return {"results":[x for x in r.json().get("results",[]) if x.get("url") and allowed(x["url"],domains,allow_any_public)],"error":""}
    except Exception: return {"results":[],"error":"search request failed"}


def add_link(paragraph, url: str):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rel)
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); props.append(color)
    under = OxmlElement("w:u"); under.set(qn("w:val"), "single"); props.append(under)
    run.append(props); text = OxmlElement("w:t"); text.text = url; run.append(text); link.append(run); paragraph._p.append(link)

@app.post("/research", dependencies=[Depends(auth)])
def research(req: ResearchRequest):
    if req.end_date < req.start_date: raise HTTPException(422, "end_date must be on/after start_date")
    regions=list(dict.fromkeys((req.regions or req.countries)))
    industries=list(dict.fromkeys(req.industry_fields))
    scope=list(dict.fromkeys((req.collection_scope or req.topics)))
    if not regions: raise HTTPException(422,"Provide at least one region, country, market, or geographic scope")
    if not industries and not scope and not req.keywords: raise HTTPException(422,"Provide an industry, collection scope, or keyword")

    # The bundled catalogue is metadata only. User-selected hosts and seed URLs
    # define source restrictions; generic tasks otherwise search public HTTPS.
    configured=[s for s in SOURCES if country_name(s["country"]) in {country_name(r) for r in regions}]
    explicit_hosts=[]
    for value in req.source_domains:
        candidate=value.strip()
        host=(urlparse(candidate if "://" in candidate else "https://"+candidate).hostname or "").lower().rstrip(".")
        if not re.fullmatch(r"(?=.{1,253}$)(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}",host):
            raise HTTPException(422,f"Invalid source domain: {value}")
        explicit_hosts.append(host)
    supplied_urls=list(dict.fromkeys([*map(str,req.urls),*map(str,req.source_urls)]))
    source_hosts=explicit_hosts
    allow_any_public=not source_hosts
    def matches(u): return allowed(u,source_hosts,allow_any_public)

    by_url={}
    # User-provided pages are processed as seeds; URLs and redirects remain HTTPS/public-host checked.
    for u in supplied_urls:
        seed_host=(urlparse(u).hostname or "").lower().rstrip(".")
        if seed_host and allowed(u,[seed_host]):
            by_url.setdefault(u,{"url":u,"query_region":regions[0],"snippet":"","user_seed":True})
    queries=build_general_queries(regions,industries,scope,req.keywords,req.start_date,req.end_date,source_hosts)
    query_log=[]
    search_enabled=bool(os.getenv("TAVILY_API_KEY"))
    if search_enabled:
        with ThreadPoolExecutor(max_workers=8) as pool:
            pending={pool.submit(search_tavily,q["query"],min(8,req.limit),q.get("domains"),allow_any_public):q for q in queries}
            completed=[]
            for future in as_completed(pending): completed.append((pending[future],future.result()))
        # Restore plan order so the audit log is stable across runs.
        completed.sort(key=lambda pair:queries.index(pair[0]))
        for q,outcome in completed:
            hits=outcome["results"]
            query_log.append({"region":q["region"],"cluster":q["cluster"],"hits":len(hits),"error":outcome["error"]})
            for hit in hits:
                u=hit.get("url","")
                if matches(u):
                    entry=by_url.setdefault(u,{"url":u,"query_region":q["region"],"snippet":hit.get("content","")})
                    if len(hit.get("content", "")) > len(entry.get("snippet", "")): entry["snippet"]=hit.get("content", "")
                    entry["search_published_date"]=hit.get("published_date","")

    stats={r:{"configured_sources":sum(country_name(s["country"])==country_name(r) for s in configured),
              "allowed_source_domains":len(source_hosts),
              "queries_run":sum(x["region"]==r for x in query_log),
              "queries_succeeded":sum(x["region"]==r and not x.get("error") for x in query_log),
              "queries_failed":sum(x["region"]==r and bool(x.get("error")) for x in query_log),
              "search_hits":0,"pages_fetched":0,
              "candidates":0,"excluded_out_of_window":0,"excluded_draft_signal":0,"excluded_by_user_terms":0} for r in regions}
    for log in query_log: stats[log["region"]]["search_hits"] += log["hits"]
    candidates=[]; excluded=[]
    # Fetch all discovered URLs before applying the per-country output cap to avoid early-result bias.
    discovered=list(by_url.items())
    def fetch_discovered(pair):
        url, found=pair
        if found.get("user_seed"):
            host=(urlparse(url).hostname or "").lower().rstrip(".")
            return fetch(url,[host])
        return fetch(url,source_hosts,allow_any_public)
    with ThreadPoolExecutor(max_workers=8) as pool:
        fetched=[(u,found,page) for (u,found),page in zip(discovered,pool.map(fetch_discovered,discovered))]
    for u, found, page in fetched:
        src=source_for_url(u, configured)
        region=country_name(src["country"]) if src else found.get("query_region", regions[0])
        if region not in stats: stats[region]={"configured_sources":0,"allowed_source_domains":len(source_hosts),"queries_run":0,"queries_succeeded":0,"queries_failed":0,"search_hits":0,"pages_fetched":0,"candidates":0,"excluded_out_of_window":0,"excluded_draft_signal":0,"excluded_by_user_terms":0}
        if page: stats[region]["pages_fetched"]+=1
        title=(page or {}).get("title") or u
        txt=(page or {}).get("text") or found.get("snippet","")
        dates=extract_published_date((page or {}).get("published_date") or found.get("search_published_date",""),(page or {}).get("structured_date",""),txt)
        d,basis=dates
        if d and not (req.start_date <= d <= req.end_date):
            stats[region]["excluded_out_of_window"]+=1
            excluded.append({"region":region,"title":title,"url":u,"reason":f"发布时间 {d.isoformat()} 不在请求区间"})
            continue
        is_draft=draft_signal(title+" "+txt)
        title_draft=draft_signal(title)
        if req.exclude_drafts and title_draft:
            stats[region]["excluded_draft_signal"]+=1
            excluded.append({"region":region,"title":title,"url":u,"reason":"草案/提案信号，待人工确认是否排除"})
            continue
        lowered=(title+" "+txt).lower()
        matched_excludes=[term for term in req.exclude_terms if term.strip().lower() in lowered]
        if matched_excludes:
            stats[region]["excluded_by_user_terms"]+=1
            excluded.append({"region":region,"title":title,"url":u,"reason":"命中用户排除词："+", ".join(matched_excludes[:5])})
            continue
        is_wechat=(urlparse(u).hostname or "").lower()=="mp.weixin.qq.com"
        topic=next((t for t in scope if t.lower() in lowered),scope[0] if len(scope)==1 else "待分类")
        industry=next((t for t in industries if t.lower() in lowered),industries[0] if len(industries)==1 else "待分类")
        item={"region":region,"country":region,"industry":industry,"title":title,"url":(page or {}).get("url",u),
              "published_date":d.isoformat() if d else "","date_basis":basis,
              "date_status":"窗口内，待核实" if d else "日期缺失或含糊，待核实",
              "topic":topic,"status":"待人工核验","draft_signal":is_draft,"draft_signal_in_title":title_draft,
              "source_level":"用户指定来源" if src is None and (not allow_any_public) else ("未分级" if src is None else ("P2公众号/二手线索" if is_wechat or src.get("level")=="P2" else src.get("level","未分级"))),
              "source_name":"微信公众号（需核验）" if is_wechat else ((src or {}).get("name",urlparse(u).hostname or "")),
              "page_fetched":bool(page),"excerpt":txt[:1200]}
        candidates.append(item)
        stats[region]["candidates"]+=1
    # Deterministic source preference; cap independently for each jurisdiction.
    rank={"P0":0,"P1":1,"P2公众号/二手线索":2,"未分级":3}
    candidates.sort(key=lambda x:(regions.index(x["region"]) if x["region"] in regions else 999,rank.get(x["source_level"],3), not bool(x["published_date"]),x["published_date"],x["title"].lower()))
    kept=[]; counts={r:0 for r in regions}
    for item in candidates:
        r=item["region"]
        if r in counts and counts[r] < req.limit: kept.append(item); counts[r]+=1
    coverage={c:{**v,"coverage_status":("候选已发现" if v["candidates"] else ("已检索但未发现候选" if v["queries_succeeded"] else ("搜索失败" if v["queries_failed"] else "未执行网络检索"))),
                 "source_status":"已配置来源约束" if v["allowed_source_domains"] else "开放域公开网页检索",
                 "note":"检索结果不等于完整覆盖；未配置或未命中的来源可能遗漏。"} for c,v in stats.items()}
    return {"task":{"task_name":req.task_name,"regions":regions,"industry_fields":industries,"collection_scope":scope,"keywords":req.keywords},
            "period":{"start":req.start_date,"end":req.end_date,"basis":"发布时间；日期筛选包含首尾日期，发现阶段前后扩展1天"},
            "search_enabled":search_enabled,"query_count":len(query_log),"searched_urls":len(by_url),"count":len(kept),
            "limit_per_region":req.limit,"coverage_by_region":coverage,"coverage_by_country":coverage,"query_log":query_log,
            "excluded_examples":excluded[:30],
            "coverage_note":"候选清单用于人工核验。日期按发布时间筛选；发布日期、来源等级和政策状态须回到原文核实。","items":kept}

@app.post("/reports", dependencies=[Depends(auth)])
def create_report(req: ReportRequest, request=None):
    if any(not item.approved for item in req.items):
        raise HTTPException(422, "Only user-approved items can be exported")
    doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Pt(48)
    normal=doc.styles["Normal"]; normal.font.name="KaiTi"; normal.font.size=Pt(12)
    normal._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia","楷体")
    doc.add_heading(req.title,0); doc.add_paragraph(f"覆盖期间：{req.period}    更新日期：{req.updated_at}")
    groups={}
    for item in req.items: groups.setdefault(item.topic or "其他",[]).append(item)
    for topic, items in groups.items():
        doc.add_heading(topic,1)
        for i,x in enumerate(items,1):
            doc.add_heading(f"{i}. {x.region or x.country}｜{x.title}",2)
            p=doc.add_paragraph(x.summary)
            doc.add_paragraph(f"发布日期：{x.published_date or '待核验'}    政策状态：{x.status}")
            p=doc.add_paragraph("相关链接：")
            for j,u in enumerate(x.sources):
                if j: p.add_run("；")
                add_link(p, str(u))
    rid=uuid.uuid4().hex; path=OUT/f"{rid}.docx"; doc.save(path)
    base=os.getenv("PUBLIC_BASE_URL", "").rstrip("/")
    return {"report_id":rid,"download_url":f"{base}/reports/{rid}/file" if base else f"/reports/{rid}/file","items":len(req.items)}

@app.get("/reports/{report_id}/file")
def download_report(report_id: str):
    path=OUT/f"{report_id}.docx"
    if not re.fullmatch(r"[a-f0-9]{32}", report_id) or not path.exists(): raise HTTPException(404,"Report not found")
    return FileResponse(path, filename="资讯收集报告.docx")
